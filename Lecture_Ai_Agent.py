# ─────────────────────────────────────────────────────────────────────────────
#  IMPORTS
# ─────────────────────────────────────────────────────────────────────────────
import os, json, re, time, traceback, html
from datetime import datetime, timedelta
from pathlib  import Path
from dateutil import parser

import pytz
import requests
import streamlit as st
from apscheduler.schedulers.background import BackgroundScheduler
from langchain_core.tools        import BaseTool
from langchain_google_genai      import ChatGoogleGenerativeAI
from langgraph.checkpoint.memory import MemorySaver
from langgraph.prebuilt          import create_react_agent
from pypdf   import PdfReader
from pptx    import Presentation
from pptx.util          import Inches, Pt
from pptx.dml.color     import RGBColor
from pptx.enum.text     import PP_ALIGN

import base64, shutil
from email.message import EmailMessage
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

# ─────────────────────────────────────────────────────────────────────────────
#  CONFIG
# ─────────────────────────────────────────────────────────────────────────────
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")
GOOGLE_CAL_KEY = os.getenv("GOOGLE_CAL_KEY")
CALENDAR_ID    = "csanduni123@gmail.com"
SL_TZ          = pytz.timezone("Asia/Colombo")
SYLLABUS_PATH  = "syllabus.pdf"
OUTPUT_DIR     = Path("output_slides"); OUTPUT_DIR.mkdir(exist_ok=True)
MEMORY_FILE    = "lecture_memory.json"
MODEL          = "models/gemini-2.5-flash"

def check_keys():
    return [k for k in ["GEMINI_API_KEY", "GOOGLE_CAL_KEY"] if not os.getenv(k)]

# ── Config ───────────────────────────────────────────────────────────────────
GEMINI_API_KEY   = os.getenv("GEMINI_API_KEY")
SL_TZ            = pytz.timezone("Asia/Colombo")
CREDENTIALS_FILE = "credentials.json"
TOKEN_FILE       = "token.json"
OUTPUT_DIR       = Path("output_slides")
MEMORY_FILE      = "lecture_memory.json"
STATUS_FILE      = "phase3_status.json"
PHASE4_STATUS    = "phase4_status.json"
MODEL            = "models/gemini-2.5-flash"
POLL_MINS        = 1
REMINDER_HOURS   = 4

CANVAS_TOKEN     = os.getenv("CANVAS_TOKEN", "7~AHDZuExYBaTxLkUhcK3KJ2BCGFERaQhAmF47CxQNkCBuRBxaH4A228EML7DX8h46")
CANVAS_COURSE    = os.getenv("CANVAS_COURSE_ID", "14583131")
CANVAS_BASE      = "https://canvas.instructure.com/api/v1"

SCOPES = [
    "https://www.googleapis.com/auth/gmail.send",
    "https://www.googleapis.com/auth/gmail.readonly",
    "https://www.googleapis.com/auth/gmail.modify",
]

APPROVE_WORDS = ["approved","approve","ok","looks good","send it",
                 "good","yes","great","perfect","fine","go ahead",
                 "confirmed","proceed","upload"]
CHANGE_WORDS  = ["change","modify","update","add","remove","slide",
                 "fix","edit","redo","replace","wrong","incorrect",
                 "revise","different","more","less","another","include",
                 "insert","delete","simplify","expand","clarify"]

OUR_MARKER = "Lecturer AI Agent — Phase 3"


# ── Auth ─────────────────────────────────────────────────────────────────────
def authenticate():
    creds = None
    if Path(TOKEN_FILE).exists():
        try:
            creds = Credentials.from_authorized_user_file(TOKEN_FILE, SCOPES)
        except Exception:
            raw   = json.load(open(TOKEN_FILE))
            creds = Credentials(
                token=raw.get("token"),
                refresh_token=raw.get("refresh_token"),
                token_uri=raw.get("token_uri","https://oauth2.googleapis.com/token"),
                client_id=raw.get("client_id"),
                client_secret=raw.get("client_secret"),
            )
    if creds and creds.expired and getattr(creds,"refresh_token",None):
        try:
            creds.refresh(Request())
            open(TOKEN_FILE,"w").write(creds.to_json())
        except Exception: pass
    if not creds:
        flow  = InstalledAppFlow.from_client_secrets_file(CREDENTIALS_FILE, SCOPES)
        creds = flow.run_local_server(port=0, prompt="consent")
        open(TOKEN_FILE,"w").write(creds.to_json())
    return creds

def get_service():
    return build("gmail","v1",credentials=authenticate())


# ── Status / Memory ───────────────────────────────────────────────────────────
def load_status():
    try:    return json.load(open(STATUS_FILE))
    except: return {}

def save_status(d):
    json.dump(d, open(STATUS_FILE,"w"), indent=2)

def load_phase3_memory():
    try:
        mem = json.load(open(MEMORY_FILE))
        return list(mem.values())[-1] if mem else {}
    except: return {}

def get_latest_pptx():
    if not OUTPUT_DIR.exists(): return None
    files = sorted(OUTPUT_DIR.glob("*.pptx"),
                   key=lambda x: x.stat().st_mtime, reverse=True)
    return files[0] if files else None


# ── Email helpers ─────────────────────────────────────────────────────────────
def send_slides_email(service, to, subject, body, pptx_path):
    msg = MIMEMultipart()
    msg["To"]=to; msg["From"]="me"; msg["Subject"]=subject
    msg.attach(MIMEText(body,"plain"))
    with open(pptx_path,"rb") as f:
        part = MIMEBase("application",
            "vnd.openxmlformats-officedocument.presentationml.presentation")
        part.set_payload(f.read())
    encoders.encode_base64(part)
    part.add_header("Content-Disposition",f'attachment; filename="{pptx_path.name}"')
    msg.attach(part)
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    r   = service.users().messages().send(userId="me",body={"raw":raw}).execute()
    return r.get("id","")

def send_plain_email(service, to, subject, body):
    msg = EmailMessage()
    msg["To"]=to; msg["From"]="me"; msg["Subject"]=subject
    msg.set_content(body)
    raw = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    r   = service.users().messages().send(userId="me",body={"raw":raw}).execute()
    return r.get("id","")


# ── Body extraction ───────────────────────────────────────────────────────────
def _extract_body(msg):
    payload = msg.get("payload",{})
    parts   = payload.get("parts",[])
    if parts:
        for part in parts:
            if part.get("mimeType")=="text/plain":
                data = part.get("body",{}).get("data","")
                if data:
                    return base64.urlsafe_b64decode(data).decode("utf-8",errors="replace").strip()
    data = payload.get("body",{}).get("data","")
    if data:
        return base64.urlsafe_b64decode(data).decode("utf-8",errors="replace").strip()
    return ""

def _strip_quoted(body):
    """Remove quoted original email — your proven clean_reply logic."""
    if not body: return ""
    if "On " in body and "wrote:" in body:
        body = body.split("On ")[0].strip()
    lines = []
    for line in body.split("\n"):
        if line.strip().startswith(">"): break
        lines.append(line)
    return "\n".join(lines).strip()

def _get_header(msg, name):
    for h in msg.get("payload",{}).get("headers",[]):
        if h["name"].lower() == name.lower():
            return h["value"]
    return ""


# ══════════════════════════════════════════════════════════════════════════════
#  ACT — send initial email
# ══════════════════════════════════════════════════════════════════════════════
def act_send_email(to_email, log=None):
    pptx = get_latest_pptx()
    if not pptx:
        raise FileNotFoundError("No .pptx in output_slides/ — run Phase 2 first.")

    info    = load_phase3_memory()
    course  = info.get("course","Machine Learning")
    week    = info.get("week","?")
    topic   = info.get("topic","this week's topic")
    subs    = info.get("subtopics",[])
    outcome = info.get("outcome","")

    now_sl  = datetime.now(SL_TZ)
    tag     = f"AGENT-REVIEW-W{week}-{now_sl.strftime('%d%m%H%M')}"
    subject = f"[{tag}] Week {week} Slides Ready — {topic} · Please Review"

    body = (
        f"Hi Dr. Perera,\n\n"
        f"Your lecture slides for Week {week} are ready for review.\n\n"
        f"Course   : {course}\n"
        f"Topic    : {topic}\n"
        f"Slides   : {pptx.name} ({pptx.stat().st_size//1024} KB · 15 slides)\n"
        f"Subtopics: {', '.join(subs)}\n"
        f"Outcome  : {outcome}\n\n"
        f"Generated: {now_sl.strftime('%A, %d %B %Y · %H:%M')} (SL time)\n\n"
        f"Please REPLY TO THIS EMAIL with:\n"
        f"  • \"Approved\" — slides uploaded to students automatically\n"
        f"  • Specific changes, e.g. \"Add examples for Classification\"\n\n"
        f"Agent checks every {POLL_MINS} min. Reminder after {REMINDER_HOURS} hrs.\n\n"
        f"{OUR_MARKER}\n"
    )

    if log: log("Authenticating Gmail…")
    service = get_service()
    if log: log(f"Sending slides to {to_email}…")
    msg_id  = send_slides_email(service, to_email, subject, body, pptx)
    if log: log(f"✓ Email sent — tracking: [{tag}]")

    status = {
        "phase":"waiting","to_email":to_email,"subject_tag":tag,
        "pptx_path":str(pptx),"course":course,"week":str(week),"topic":topic,
        "sent_at":now_sl.isoformat(),"sent_date":now_sl.strftime("%Y/%m/%d"),
        "reminder_sent":False,"revision_count":0,"reply_body":"",
        "msg_id":msg_id,"last_polled":now_sl.isoformat(),
    }
    save_status(status)
    return {"sent":True,"subject":subject,"pptx":pptx.name,"tag":tag}


# ══════════════════════════════════════════════════════════════════════════════
#  OBSERVE — poll Gmail
#  FIX A: checks subject for "Re:" to find replies reliably
# ══════════════════════════════════════════════════════════════════════════════
def observe_poll(log=None):
    status = load_status()
    if not status: return {"result":"no_status"}

    tag       = status.get("subject_tag","")
    sent_date = status.get("sent_date","")
    sent_at   = datetime.fromisoformat(status["sent_at"])
    now_sl    = datetime.now(SL_TZ)
    elapsed_h = (now_sl - sent_at).total_seconds() / 3600

    if log: log(f"Polling [{tag}]… ({elapsed_h:.1f} hrs elapsed)")

    service  = get_service()
    query    = f'subject:"{tag}" after:{sent_date}'
    try:
        results  = service.users().messages().list(
            userId="me", q=query, maxResults=10).execute()
        messages = results.get("messages",[])
    except Exception as e:
        if log: log(f"Gmail search error: {e}")
        return {"result":"error","msg":str(e)}

    if log: log(f"Found {len(messages)} message(s) with tag")

    # FIX A — check "Re:" in subject to identify replies
    reply_body = None
    for m_ref in messages:
        msg     = service.users().messages().get(
            userId="me", id=m_ref["id"], format="full").execute()
        subject = _get_header(msg,"Subject")
        sender  = _get_header(msg,"From")
        if log: log(f"  Subject: {subject[:60]} | From: {sender[:35]}")

        if "re:" not in subject.lower():
            if log: log("  → skipping (no Re: — this is our sent email)")
            continue

        # This is a reply — extract and clean
        raw_body   = _extract_body(msg)
        reply_body = _strip_quoted(raw_body)
        if log: log(f"  → reply found: '{reply_body[:80]}'")
        break

    if reply_body is not None:
        b = reply_body.lower()

        # Check CHANGE words first — before approval
        # (prevents "Good, but change slide 4" from triggering approval)
        if any(w in b for w in CHANGE_WORDS):
            status["phase"]      = "changes"
            status["reply_body"] = reply_body
            save_status(status)
            if log: log("📝 Changes requested — triggering revision loop")
            return {"result":"changes","body":reply_body}

        if any(w in b for w in APPROVE_WORDS):
            # Save approved pptx with _APPROVED suffix for Phase 4
            pptx_path    = Path(status.get("pptx_path",""))
            approved_path = OUTPUT_DIR / f"{pptx_path.stem}_APPROVED.pptx"
            if pptx_path.exists():
                shutil.copy(pptx_path, approved_path)
                if log: log(f"✓ Approved file saved: {approved_path.name}")

            status["phase"]         = "approved"
            status["reply_body"]    = reply_body
            status["approved_at"]   = now_sl.isoformat()
            status["approved_pptx"] = str(approved_path)
            save_status(status)
            if log: log("✅ APPROVED!")
            return {"result":"approved","body":reply_body}

        # Unclear reply
        status["phase"]      = "unclear"
        status["reply_body"] = reply_body
        save_status(status)
        if log: log(f"Reply unclear: '{reply_body[:100]}'")
        return {"result":"unclear","body":reply_body}

    # No reply yet
    if log: log("No lecturer reply yet.")

    if elapsed_h >= REMINDER_HOURS and not status.get("reminder_sent"):
        if log: log(f"⏰ {REMINDER_HOURS} hrs passed — sending reminder…")
        send_plain_email(
            service, status["to_email"],
            f"[REMINDER] [{tag}] Please review Week {status['week']} slides",
            f"Hi Dr. Perera,\n\nWeek {status['week']} slides have been waiting "
            f"{elapsed_h:.0f} hours for review.\nPlease reply 'Approved' or changes.\n\n"
            f"Lecturer AI Agent"
        )
        status["reminder_sent"] = True
        save_status(status)
        if log: log("✓ Reminder sent")
        return {"result":"reminder_sent","elapsed_h":round(elapsed_h,1)}

    status["last_polled"] = now_sl.isoformat()
    save_status(status)
    return {"result":"waiting","elapsed_h":round(elapsed_h,1)}


# ══════════════════════════════════════════════════════════════════════════════
#  THINK — parse change instructions
# ══════════════════════════════════════════════════════════════════════════════
def think_parse_changes(reply_body, log=None):
    if log: log("Think — Gemini parsing change instructions…")

    # Regex fallback: extract numbered list items
    numbered = re.findall(r"\d+[\.\)]\s*(.+?)(?=\n\d+[\.\)]|\Z)", reply_body, re.DOTALL)
    numbered = [n.strip() for n in numbered if n.strip()]

    # Extract any slide numbers mentioned
    slide_nums = list(set(int(n) for n in re.findall(r"[Ss]lide\s*(\d+)", reply_body)))

    # Detect special flags
    b = reply_body.lower()
    has_diagram  = any(w in b for w in ["diagram","visual","figure","illustration","chart"])
    has_example  = any(w in b for w in ["example","examples","worked","demonstrate","show"])
    has_simplify = any(w in b for w in ["simplify","simple","easier","undergrad","basic","too technical"])

    if not GEMINI_API_KEY:
        instr = numbered if numbered else [reply_body]
        return {"slides":slide_nums,"instructions":instr,
                "general":" | ".join(instr[:3]),"raw":reply_body,
                "has_diagram":has_diagram,"has_example":has_example,
                "has_simplify":has_simplify}

    llm = ChatGoogleGenerativeAI(model=MODEL,google_api_key=GEMINI_API_KEY,temperature=0)
    try:
        resp = llm.invoke(
            f'Lecturer change request for lecture slides:\n"{reply_body}"\n\n'
            'Return ONLY this JSON (no markdown):\n'
            '{"slides":[list of slide numbers mentioned, or []],'
            '"instructions":["each change as one clear sentence"],'
            '"general":"one sentence summary of all changes combined"}'
        )
        raw  = resp.content if hasattr(resp,"content") else str(resp)
        raw  = re.sub(r"```json\s*|\s*```","",raw).strip()
        data = json.loads(raw)
        data["raw"]         = reply_body
        data["has_diagram"] = has_diagram
        data["has_example"] = has_example
        data["has_simplify"]= has_simplify
        # Merge slide nums
        existing = data.get("slides",[])
        data["slides"] = list(set(existing + slide_nums))
        n = len(data.get("instructions",[]))
        s = data.get("slides",[])
        if log: log(f"✓ Think: {n} instructions, slides {s}")
        if log: log(f"  Summary: {data.get('general','')}")
        return data
    except Exception as e:
        if log: log(f"Gemini parse failed ({e}) — using regex")
        instr = numbered if numbered else [reply_body]
        return {"slides":slide_nums,"instructions":instr,
                "general":" | ".join(instr[:3]),"raw":reply_body,
                "has_diagram":has_diagram,"has_example":has_example,
                "has_simplify":has_simplify}


# ══════════════════════════════════════════════════════════════════════════════
#  PPTX BUILDER — same engine as Phase 2
#  FIX B: fully rebuilds slides from scratch, not just patching text boxes
# ══════════════════════════════════════════════════════════════════════════════
def build_pptx_phase3(slides_data, out_path, course, week, topic):
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    NA=RGBColor(0x1e,0x1b,0x4b); IN=RGBColor(0x63,0x66,0xf1)
    BL=RGBColor(0x37,0x8a,0xdd); GO=RGBColor(0xea,0xb3,0x08)
    WH=RGBColor(0xff,0xff,0xff); LT=RGBColor(0xf8,0xfa,0xff)
    GR=RGBColor(0x6b,0x72,0x80); LA=RGBColor(0xa5,0xb4,0xfc)
    MI=RGBColor(0xee,0xf2,0xff); GN=RGBColor(0x05,0x96,0x69)
    RE=RGBColor(0xdc,0x26,0x26); DI=RGBColor(0xc7,0xd2,0xfe)
    layout = prs.slide_layouts[6]

    def R(sl,l,t,w,h,c):
        s=sl.shapes.add_shape(1,Inches(l),Inches(t),Inches(w),Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb=c; s.line.fill.background()
    def T(sl,tx,l,t,w,h,sz=18,bold=False,c=None,al=PP_ALIGN.LEFT):
        tb=sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
        tf=tb.text_frame; tf.word_wrap=True; p=tf.paragraphs[0]; p.alignment=al
        r=p.add_run(); r.text=str(tx); r.font.size=Pt(sz); r.font.bold=bold
        if c: r.font.color.rgb=c
    def B(sl,items,l,t,w,h,sz=17,tc=None,dc=None):
        if not items: return
        tb=sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
        tf=tb.text_frame; tf.word_wrap=True; first=True
        for item in items:
            item=str(item).strip()
            if not item: continue
            p=tf.paragraphs[0] if first else tf.add_paragraph(); first=False
            p.alignment=PP_ALIGN.LEFT; p.space_after=Pt(7)
            d=p.add_run(); d.text="▸  "; d.font.size=Pt(sz-2); d.font.color.rgb=dc or IN
            r=p.add_run(); r.text=item; r.font.size=Pt(sz); r.font.color.rgb=tc or NA

    for idx,s in enumerate(slides_data[:15]):
        sl    = prs.slides.add_slide(layout)
        stype = s.get("type","content")
        stitle= s.get("title",f"Slide {idx+1}")
        items = [b.strip() for b in s.get("content","").split("|") if b.strip()]
        snum  = idx+1

        if snum==1:
            R(sl,0,0,13.33,7.5,NA); R(sl,0,5.8,13.33,1.7,IN); R(sl,0.6,3.35,5.5,0.06,GO)
            for i,x in enumerate([11.5,11.9,12.3]):
                sh=sl.shapes.add_shape(9,Inches(x),Inches(0.35),Inches(0.25),Inches(0.25))
                sh.fill.solid(); sh.fill.fore_color.rgb=[IN,GO,BL][i]; sh.line.fill.background()
            T(sl,course,0.6,1.5,12,0.5,sz=15,c=LA)
            T(sl,topic,0.6,2.3,11.5,1.8,sz=38,bold=True,c=WH)
            T(sl,f"Week {week}  ·  Lecture Notes",0.6,5.95,8,0.5,sz=13,c=WH)
            T(sl,datetime.now(SL_TZ).strftime("%B %Y"),9.5,5.95,3.5,0.5,sz=13,c=WH,al=PP_ALIGN.RIGHT)
        else:
            R(sl,0,0,13.33,7.5,LT); R(sl,0,0,13.33,1.25,NA); R(sl,0,0,0.1,7.5,IN)
            R(sl,11.8,0.38,1.35,0.42,IN)
            T(sl,f"{snum}/15",11.82,0.39,1.3,0.4,sz=11,bold=True,c=WH,al=PP_ALIGN.CENTER)
            T(sl,f"{course}  ·  Week {week}",0.25,0.06,9,0.38,sz=11,c=LA)
            T(sl,stitle,0.25,0.42,11.3,0.7,sz=24,bold=True,c=WH)
            if stype=="summary":
                R(sl,0.25,1.35,12.8,5.85,MI); B(sl,items,0.55,1.5,12.3,5.6,sz=20,tc=NA,dc=GN)
            elif stype=="example":
                R(sl,0.25,1.35,0.06,5.85,GO); T(sl,"Worked Example",0.45,1.38,6,0.45,sz=13,c=GR)
                B(sl,items,0.55,1.9,12.3,5,sz=18,dc=GO)
            elif stype=="diagram":
                R(sl,0.25,1.35,0.06,5.85,IN); T(sl,"Visual Diagram",0.45,1.38,6,0.45,sz=13,c=LA)
                B(sl,items,0.55,1.9,12.3,5,sz=18,dc=IN)
            elif stype=="tips":
                R(sl,0.25,1.35,0.06,5.85,RE); B(sl,items,0.55,1.5,12.3,5.6,sz=18,dc=RE)
            elif len(items)>=6:
                mid=len(items)//2; B(sl,items[:mid],0.25,1.4,6.2,5.8,sz=17)
                R(sl,6.6,1.45,0.03,5.6,DI); B(sl,items[mid:],6.7,1.4,6.3,5.8,sz=17)
            else:
                B(sl,items,0.35,1.45,12.6,5.8,sz=19)
            R(sl,0.1,7.15,13.2,0.02,DI)

    while len(prs.slides)<15:
        sl=prs.slides.add_slide(layout); n=len(prs.slides)
        R(sl,0,0,13.33,7.5,LT); R(sl,0,0,13.33,1.25,NA); R(sl,0,0,0.1,7.5,IN)
        T(sl,f"Slide {n}",0.25,0.42,11.3,0.7,sz=24,bold=True,c=WH)
    prs.save(str(out_path))


# ══════════════════════════════════════════════════════════════════════════════
#  ACT — regenerate revised slides and re-send
#  FIX B: fully rebuilds all 15 slides with changes applied throughout
# ══════════════════════════════════════════════════════════════════════════════
def act_regenerate_and_resend(log=None):
    status   = load_status()
    old_path = Path(status.get("pptx_path",""))
    reply    = status.get("reply_body","")
    course   = status.get("course","Course")
    topic    = status.get("topic","Topic")
    week     = status.get("week","?")
    to_email = status.get("to_email","")
    rev_num  = status.get("revision_count",0) + 1
    subs     = load_phase3_memory().get("subtopics",[])
    outcome  = load_phase3_memory().get("outcome","")

    # Think — parse what needs changing
    changes = think_parse_changes(reply, log=log)

    inst_list = changes.get("instructions",[reply])
    inst_text = "\n".join(f"  {i+1}. {ins}" for i,ins in enumerate(inst_list))
    sns       = changes.get("slides",[])
    general   = changes.get("general", reply[:100])

    # New version filename
    ver      = rev_num + 1
    new_name = re.sub(r"_v\d+(_APPROVED)?\.pptx$","",old_path.stem) + f"_v{ver}.pptx"
    new_path = OUTPUT_DIR / new_name

    if log: log(f"Act — generating revised slides → {new_name}")
    if log: log(f"  Changes: {general}")

    # FIX B — ask Gemini to regenerate ALL 15 slides incorporating the changes
    if GEMINI_API_KEY:
        time.sleep(5)
        llm = ChatGoogleGenerativeAI(
            model=MODEL, google_api_key=GEMINI_API_KEY, temperature=0.4)

        if log: log("Calling Gemini to regenerate all 15 slides with changes applied…")

        extra = ""
        if changes.get("has_diagram"):  extra += "\n- Include a dedicated diagram/visual description slide."
        if changes.get("has_example"):  extra += "\n- Add a worked example slide with step-by-step solutions."
        if changes.get("has_simplify"): extra += "\n- Use simpler language suitable for undergraduates throughout."

        prompt = f"""
You are regenerating a 15-slide university lecture presentation with lecturer feedback applied.

Course   : {course}
Week     : {week}
Topic    : {topic}
Subtopics: {', '.join(subs)}
Outcome  : {outcome}

Lecturer change instructions:
{inst_text}
{f"Focus especially on slides: {sns}" if sns else "Apply changes throughout all slides."}
{extra}

Generate ALL 15 slides with these changes fully incorporated.
Return ONLY a valid JSON array of 15 slide objects. No markdown. No backticks.
Each: {{"slide":N,"type":"string","title":"string","content":"point1|point2|point3|point4"}}
Types: title, overview, content, example, diagram, summary, tips, next
Use | between bullets. Max 8 words per bullet. Real accurate content about {topic}.
Make sure every instruction above is visibly applied in the slides.
Return ONLY the JSON array.
"""
        try:
            resp    = llm.invoke(prompt)
            raw_out = resp.content if hasattr(resp,"content") else str(resp)
            raw_out = re.sub(r"```json\s*|\s*```","",raw_out).strip()
            try:
                slides_data = json.loads(raw_out)
            except Exception:
                m = re.search(r"\[.*\]", raw_out, re.DOTALL)
                slides_data = json.loads(m.group()) if m else []
            if log: log(f"✓ Gemini returned {len(slides_data)} slides")
        except Exception as e:
            if log: log(f"Gemini error: {e} — copying original file")
            shutil.copy(old_path, new_path)
            slides_data = []

        if slides_data:
            if log: log(f"Building pptx: {new_name}")
            build_pptx_phase3(slides_data, new_path, course, week, topic)
            if log: log(f"✓ {new_path.name} saved ({new_path.stat().st_size//1024} KB)")
    else:
        shutil.copy(old_path, new_path)
        if log: log("No Gemini key — copied original file")

    # Re-send email with new tracking tag
    now_sl  = datetime.now(SL_TZ)
    tag     = f"AGENT-REVIEW-W{week}-REV{rev_num}-{now_sl.strftime('%d%m%H%M')}"
    subject = f"[{tag}] Week {week} Slides REVISED v{ver} — {topic}"
    inst_display = "\n".join(f"  ✓ {ins}" for ins in inst_list)
    body = (
        f"Hi Dr. Perera,\n\n"
        f"The slides have been revised based on your feedback.\n\n"
        f"Changes applied:\n{inst_display}\n\n"
        f"File     : {new_path.name}\n"
        f"Revision : #{rev_num} (version {ver})\n"
        f"Generated: {now_sl.strftime('%A, %d %B %Y · %H:%M')} (SL time)\n\n"
        f"Please REPLY TO THIS EMAIL with:\n"
        f"  • \"Approved\" — slides sent to students automatically\n"
        f"  • Further changes if needed\n\n"
        f"{OUR_MARKER}\n"
    )

    if log: log(f"Re-sending to {to_email}…")
    service = get_service()
    msg_id  = send_slides_email(service, to_email, subject, body, new_path)
    if log: log(f"✓ Revision {rev_num} email sent — tag: [{tag}]")

    status.update({
        "phase":"waiting","subject_tag":tag,"pptx_path":str(new_path),
        "sent_at":now_sl.isoformat(),"sent_date":now_sl.strftime("%Y/%m/%d"),
        "reminder_sent":False,"revision_count":rev_num,
        "reply_body":"","msg_id":msg_id,"last_polled":now_sl.isoformat(),
    })
    save_status(status)
    return {"revision":rev_num,"new_pptx":new_path.name,"general":general}


# ══════════════════════════════════════════════════════════════════════════════
#  PHASE 4 — Canvas LMS Upload + Student Notification
#  Logic preserved from canvas.py; UI is integrated into existing panels only.
# ══════════════════════════════════════════════════════════════════════════════
def load_phase4_status() -> dict:
    try:    return json.load(open(PHASE4_STATUS))
    except: return {}

def save_phase4_status(d: dict):
    with open(PHASE4_STATUS, "w") as f:
        json.dump(d, f, indent=2)

def canvas_headers() -> dict:
    return {"Authorization": f"Bearer {CANVAS_TOKEN}"}

def load_canvas_memory() -> dict:
    try:
        mem = json.load(open(MEMORY_FILE))
        return list(mem.values())[-1] if mem else {}
    except: return {}

def get_approved_pptx() -> Path | None:
    p3 = load_status()
    approved = p3.get("approved_pptx", "")
    if approved and Path(approved).exists():
        return Path(approved)

    if OUTPUT_DIR.exists():
        approved_files = sorted(
            OUTPUT_DIR.glob("*_APPROVED.pptx"),
            key=lambda x: x.stat().st_mtime, reverse=True)
        if approved_files:
            return approved_files[0]

    if OUTPUT_DIR.exists():
        all_files = sorted(
            OUTPUT_DIR.glob("*.pptx"),
            key=lambda x: x.stat().st_mtime, reverse=True)
        if all_files:
            return all_files[0]

    return None

def upload_to_canvas(pptx_path: Path, week: str, topic: str, log=None) -> dict:
    folder_path = f"Week {week} — {topic}"[:50]
    now_sl      = datetime.now(SL_TZ)

    if log: log("Phase 4 — Requesting Canvas upload URL")

    url  = f"{CANVAS_BASE}/courses/{CANVAS_COURSE}/files"
    data = {
        "name":                pptx_path.name,
        "size":                pptx_path.stat().st_size,
        "content_type":        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "parent_folder_path":  folder_path,
    }
    res = requests.post(url, headers=canvas_headers(), data=data, timeout=30)

    if res.status_code not in (200, 201):
        raise RuntimeError(
            f"Canvas upload request failed: {res.status_code} — {res.text[:300]}")

    upload_data   = res.json()
    upload_url    = upload_data.get("upload_url")
    upload_params = upload_data.get("upload_params", {})

    if not upload_url:
        raise RuntimeError(f"No upload_url in Canvas response: {upload_data}")

    if log: log("Phase 4 — Uploading approved slides to Canvas")
    with open(pptx_path, "rb") as f:
        upload_res = requests.post(
            upload_url,
            data=upload_params,
            files={"file": (pptx_path.name, f,
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
            timeout=60,
        )

    if upload_res.status_code in (200, 201):
        try:
            file_info = upload_res.json()
        except Exception:
            file_info = {}
    elif upload_res.status_code in (301, 302, 303):
        confirm_url = upload_res.headers.get("Location", "")
        if confirm_url:
            confirm_res = requests.get(
                confirm_url, headers=canvas_headers(), timeout=30)
            file_info   = confirm_res.json() if confirm_res.ok else {}
        else:
            file_info = {}
    else:
        raise RuntimeError(
            f"Canvas file upload failed: {upload_res.status_code} — {upload_res.text[:300]}")

    file_id   = file_info.get("id", "unknown")
    file_name = file_info.get("display_name", pptx_path.name)
    file_url  = file_info.get("url", "")

    if log: log(f"Canvas upload complete — file ID {file_id}")

    return {
        "file_id":   file_id,
        "file_name": file_name,
        "file_url":  file_url,
        "folder":    folder_path,
        "uploaded_at": now_sl.isoformat(),
    }

def post_announcement(week: str, topic: str, course: str, file_info: dict, log=None) -> dict:
    now_sl    = datetime.now(SL_TZ)
    file_name = file_info.get("file_name", "slides.pptx")
    folder    = file_info.get("folder", f"Week {week}")

    title = f"Week {week} Lecture Slides Available — {topic}"

    message = f"""
<p>Dear Students,</p>

<p>The lecture slides for <strong>Week {week}: {topic}</strong> are now 
available in Canvas.</p>

<table style="border-collapse:collapse;width:100%;max-width:500px">
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">Course</td>
      <td style="padding:6px 12px">{course}</td></tr>
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">Week</td>
      <td style="padding:6px 12px">Week {week}</td></tr>
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">Topic</td>
      <td style="padding:6px 12px">{topic}</td></tr>
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">File</td>
      <td style="padding:6px 12px">{file_name}</td></tr>
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">Folder</td>
      <td style="padding:6px 12px">{folder}</td></tr>
  <tr><td style="padding:6px 12px;background:#f0f4ff;font-weight:bold">Posted</td>
      <td style="padding:6px 12px">{now_sl.strftime("%A, %d %B %Y · %H:%M")} (SL)</td></tr>
</table>

<p>You can find the slides in the <strong>Files</strong> section under 
<em>{folder}</em>.</p>

<p>Please review the slides before the lecture. If you have any questions, 
contact your lecturer.</p>

<p><em>— Lecturer AI Agent (automated notification)</em></p>
"""

    if log: log("Phase 4 — Posting Canvas announcement")

    url  = f"{CANVAS_BASE}/courses/{CANVAS_COURSE}/discussion_topics"
    data = {
        "title":         title,
        "message":       message,
        "is_announcement": True,
        "published":     True,
    }
    res = requests.post(url, headers=canvas_headers(), data=data, timeout=30)

    if res.status_code in (200, 201):
        ann  = res.json()
        ann_id  = ann.get("id", "unknown")
        ann_url = ann.get("html_url", "")
        if log: log(f"Canvas announcement posted — ID {ann_id}")
        return {"announcement_id": ann_id, "announcement_url": ann_url, "title": title}
    else:
        if log: log("Canvas announcement failed — file upload still complete")
        return {"announcement_id": None, "error": res.text[:200]}

def verify_canvas(log=None) -> dict:
    try:
        res = requests.get(
            f"{CANVAS_BASE}/courses/{CANVAS_COURSE}",
            headers=canvas_headers(), timeout=10)
        if res.ok:
            info = res.json()
            name = info.get("name", "Unknown")
            if log: log(f"Canvas connected — Course: {name}")
            return {"ok": True, "course_name": name}
        else:
            return {"ok": False, "error": f"HTTP {res.status_code}: {res.text[:200]}"}
    except Exception as e:
        return {"ok": False, "error": str(e)}

def run_phase4(log=None) -> dict:
    p3     = load_status()
    memory = load_canvas_memory()
    week   = str(p3.get("week",   memory.get("week",   "?")))
    topic  = p3.get("topic",      memory.get("topic",  "Lecture"))
    course = p3.get("course",     memory.get("course", "Course"))

    if p3.get("phase") != "approved":
        raise RuntimeError(
            "Phase 3 not approved yet. Slides must be approved by lecturer before uploading.")

    pptx = get_approved_pptx()
    if not pptx:
        raise FileNotFoundError(
            "No approved .pptx found. Run Phase 3 first and get lecturer approval.")

    if log: log(f"Phase 4 — Found approved file {pptx.name}")

    file_info = upload_to_canvas(pptx, week, topic, log=log)
    ann_info = post_announcement(week, topic, course, file_info, log=log)

    now_sl = datetime.now(SL_TZ)
    result = {
        "phase":            "complete",
        "pptx_uploaded":    pptx.name,
        "canvas_file_id":   file_info.get("file_id"),
        "canvas_folder":    file_info.get("folder"),
        "announcement_id":  ann_info.get("announcement_id"),
        "announcement_url": ann_info.get("announcement_url",""),
        "uploaded_at":      now_sl.isoformat(),
        "week":             week,
        "topic":            topic,
        "course":           course,
    }
    save_phase4_status(result)

    summary = (
        f"Phase 4 complete. "
        f"Uploaded {pptx.name} to Canvas folder {file_info.get('folder')} · "
        f"students notified."
    )
    if log: log(summary)
    return {**result, "summary": summary}

def run_phase4_once(log=None) -> dict:
    previous = load_phase4_status()
    approved = get_approved_pptx()
    approved_name = approved.name if approved else ""
    if previous.get("phase") == "complete" and previous.get("pptx_uploaded") == approved_name:
        if log: log("Phase 4 already complete — Canvas upload confirmed")
        return previous
    return run_phase4(log=log)




# ─────────────────────────────────────────────────────────────────────────────
#  PHASE 1 — TOOLS
# ─────────────────────────────────────────────────────────────────────────────
class GoogleCalendarTool(BaseTool):
    name: str        = "google_calendar_fetcher"
    description: str = (
        "Fetches today's and tomorrow's lecture schedule from Google Calendar. "
        "Skips past events and already-ended classes. "
        "Returns JSON list with summary, start, end, location, urgency, "
        "hours_until, day_label. Call this tool first."
    )

    def _run(self, query: str = "") -> str:
        url = (
            f"https://www.googleapis.com/calendar/v3/calendars/"
            f"{CALENDAR_ID}/events?key={GOOGLE_CAL_KEY}&timeZone=Asia/Colombo"
        )
        try:
            resp = requests.get(url, timeout=10)
            resp.raise_for_status()
            raw = resp.json().get("items", [])
        except Exception as e:
            return json.dumps({"error": str(e)})

        now_sl      = datetime.now(SL_TZ)
        today_sl    = now_sl.date()
        tomorrow_sl = today_sl + timedelta(days=1)
        parsed      = []

        for event in raw:
            s_str = event["start"].get("dateTime", event["start"].get("date", ""))
            e_str = event["end"].get("dateTime",   event["end"].get("date",   ""))
            try:
                s_dt = parser.isoparse(s_str).astimezone(SL_TZ)
                e_dt = parser.isoparse(e_str).astimezone(SL_TZ)
            except Exception:
                continue

            if s_dt.date() not in (today_sl, tomorrow_sl):
                continue
            if e_dt < now_sl:
                continue

            hrs       = (s_dt - now_sl).total_seconds() / 3600
            day_label = "today" if s_dt.date() == today_sl else "tomorrow"
            urgency   = (
                "ongoing"  if hrs < 0  else
                "critical" if hrs < 3  else
                "soon"     if hrs < 6  else
                "planned"  if day_label == "today" else
                "tomorrow"
            )

            parsed.append({
                "summary":     event.get("summary",  "No title"),
                "location":    event.get("location", "No location"),
                "start":       s_dt.strftime("%Y-%m-%d %H:%M"),
                "end":         e_dt.strftime("%Y-%m-%d %H:%M"),
                "hours_until": round(hrs, 2),
                "urgency":     urgency,
                "day_label":   day_label,
            })

        return json.dumps(parsed, indent=2)

    def _arun(self, q=""):
        raise NotImplementedError


class ScheduleAnalyserTool(BaseTool):
    name: str        = "schedule_analyser"
    description: str = (
        "Analyses JSON from google_calendar_fetcher. Detects busy days. "
        "Builds slide-prep queue — EXCLUDES ongoing classes since slides "
        "cannot be prepared for a class already in session. "
        "Saves priority_queue.json and all_events.json."
    )

    def _run(self, events_json: str) -> str:
        try:
            events = json.loads(events_json)
        except json.JSONDecodeError:
            return "Error: invalid JSON. Run google_calendar_fetcher first."

        if not events:
            return "No upcoming events found for today or tomorrow."

        day_counts: dict = {}
        for e in events:
            d = e.get("day_label", "unknown")
            day_counts[d] = day_counts.get(d, 0) + 1

        rank  = {"critical": 1, "soon": 2, "planned": 3, "tomorrow": 4}
        queue = []

        for e in events:
            urg       = e.get("urgency", "planned")
            day_label = e.get("day_label", "today")
            is_busy   = day_counts.get(day_label, 0) >= 2

            if urg == "ongoing":
                continue

            queue.append({
                **e,
                "urgency_rank": rank.get(urg, 5),
                "busy_day":     is_busy,
                "action": (
                    "prepare_slides_NOW"     if urg == "critical"             else
                    "prepare_slides_soon"    if urg in ("soon", "planned")    else
                    "prepare_slides_tonight"
                ),
            })

        queue.sort(key=lambda x: (x["urgency_rank"], x["start"]))

        with open("priority_queue.json", "w") as f:
            json.dump(queue, f, indent=2)
        with open("all_events.json", "w") as f:
            json.dump(events, f, indent=2)

        busy_days = [d for d, c in day_counts.items() if c >= 2]
        top       = queue[0] if queue else None
        top_text  = (
            f"'{top['summary']}' at {top['start'][11:]} [{top['urgency']}]"
            if top else "none (all classes ongoing or ended)"
        )
        return (
            f"Analysis complete.\n"
            f"• Classes today: {sum(1 for e in events if e['day_label']=='today')}\n"
            f"• Classes tomorrow: {sum(1 for e in events if e['day_label']=='tomorrow')}\n"
            f"• Busy days: {', '.join(busy_days) if busy_days else 'none'}\n"
            f"• Slide prep queue: {len(queue)} lectures\n"
            f"• Most urgent for slides: {top_text}\n"
            f"• File saved: priority_queue.json ✓"
        )

    def _arun(self, q=""):
        raise NotImplementedError


# ─────────────────────────────────────────────────────────────────────────────
#  PHASE 1 — AGENT
# ─────────────────────────────────────────────────────────────────────────────
def build_agent():
    llm    = ChatGoogleGenerativeAI(
        model=MODEL, google_api_key=GEMINI_API_KEY, temperature=0)
    tools  = [GoogleCalendarTool(), ScheduleAnalyserTool()]
    memory = MemorySaver()
    return create_react_agent(model=llm, tools=tools, checkpointer=memory)


AGENT_PROMPT = """
You are an autonomous lecturer assistant. Do exactly these two steps:
1. Call google_calendar_fetcher to get today's and tomorrow's lectures.
2. Pass its full JSON output to schedule_analyser.
Return only the plain text summary that schedule_analyser gives you. Nothing else.
"""


def extract_clean_text(result: dict) -> str:
    for msg in reversed(result.get("messages", [])):
        content = getattr(msg, "content", None)
        if not content:
            continue
        if isinstance(content, str) and content.strip():
            return content.strip()
        if isinstance(content, list):
            for block in content:
                if isinstance(block, dict) and block.get("type") == "text":
                    t = block.get("text", "").strip()
                    if t:
                        return t
    return "Agent completed. Check priority_queue.json for results."


def run_phase1_agent() -> str:
    agent  = build_agent()
    config = {"configurable": {"thread_id": "phase1"}}
    result = agent.invoke(
        {"messages": [{"role": "user", "content": AGENT_PROMPT}]},
        config=config,
    )
    return extract_clean_text(result)


# ─────────────────────────────────────────────────────────────────────────────
#  SCHEDULER (7 AM auto-run)
# ─────────────────────────────────────────────────────────────────────────────
def run_agent_job():
    try:
        result = run_phase1_agent()
        with open("last_agent_run.json", "w") as f:
            json.dump({"time": datetime.now(SL_TZ).strftime("%Y-%m-%d %H:%M:%S"),
                       "result": result, "mode": "auto-7AM"}, f, indent=2)
    except Exception as e:
        with open("last_agent_run.json", "w") as f:
            json.dump({"error": str(e)}, f)


def start_scheduler():
    s = BackgroundScheduler(timezone=SL_TZ)
    s.add_job(run_agent_job, trigger="cron", hour=7, minute=0, id="phase1")
    s.start()


if "scheduler_started" not in st.session_state:
    start_scheduler()
    st.session_state["scheduler_started"] = True


# ─────────────────────────────────────────────────────────────────────────────
#  PHASE 2 — PIPELINE
# ─────────────────────────────────────────────────────────────────────────────
def load_memory():
    try:    return json.load(open(MEMORY_FILE))
    except: return {}


def save_memory(week, course, topic, subtopics, outcome, nf, sf):
    mem = load_memory()
    mem[f"week{week}"] = {
        "week": week, "course": course, "topic": topic,
        "subtopics": subtopics, "outcome": outcome,
        "notes_file": nf, "slides_file": sf,
        "generated": datetime.now(SL_TZ).strftime("%Y-%m-%d %H:%M:%S"),
    }
    json.dump(mem, open(MEMORY_FILE, "w"), indent=2)


def parse_syllabus(path):
    text = "\n".join(p.extract_text() or "" for p in PdfReader(path).pages)

    def field(k):
        m = re.search(rf"{k}:\s*(.+)", text)
        return m.group(1).strip() if m else ""

    cn = field("COURSE_NAME"); cc = field("COURSE_CODE"); sd = field("START_DATE")
    try:    start = datetime.strptime(sd, "%Y-%m-%d").date()
    except: start = None
    today = datetime.now(SL_TZ).date()
    wn    = max(1, ((today - start).days // 7) + 1) if start else 1
    blocks = re.findall(
        r"WEEK:\s*(\d+)\s*[\r\n]+TITLE:\s*(.+?)[\r\n]+TOPICS:\s*[\r\n]+(.*?)(?=WEEK:\s*\d|\Z)",
        text, re.DOTALL)
    sched = {}
    for w, title, raw in blocks:
        topics  = [t.strip().lstrip("- ") for t in raw.split("\n") if t.strip().startswith("-")]
        om      = re.search(r"OUTCOME:\s*(.+)", raw)
        outcome = om.group(1).strip() if om else ""
        sched[int(w)] = {"title": title.strip(), "topics": topics, "outcome": outcome}
    cur = sched.get(wn, {})
    if not cur and sched:
        wn = max(sched.keys()); cur = sched[wn]
    return {
        "course_name": cn, "course_code": cc,
        "start_date": str(start) if start else "", "today": str(today),
        "week_num": wn, "topic": cur.get("title", ""),
        "subtopics": cur.get("topics", []), "outcome": cur.get("outcome", ""),
        "full_schedule": {str(k): v for k, v in sched.items()},
    }


def gemini_call(prompt, temp=0.3):
    llm = ChatGoogleGenerativeAI(model=MODEL, google_api_key=GEMINI_API_KEY, temperature=temp)
    r   = llm.invoke(prompt)
    return r.content if hasattr(r, "content") else str(r)


def generate_notes(info, log=None):
    cn   = f"{info['course_name']} ({info['course_code']})"
    wn   = info["week_num"]; tp = info["topic"]
    subs = "\n".join(f"  - {s}" for s in info["subtopics"])
    if log: log(f"Gemini call 1 — lecture notes for {tp}…")
    notes = gemini_call(
        f"You are an expert lecturer for {cn}.\n"
        f"Generate structured notes for Week {wn}: {tp}.\n"
        f"Subtopics:\n{subs}\nOutcome: {info['outcome']}\n\n"
        f"Structure: TOPIC / OVERVIEW / DEFINITIONS / KEY CONCEPTS / "
        f"EXAMPLES / DIAGRAM DESCRIPTIONS / SUMMARY / EXAM TIPS\n"
        f"Be clear and accurate for undergraduates.", 0.3)
    fname = f"notes_week{wn}_{tp.replace(' ','_')[:20]}.txt"
    (OUTPUT_DIR / fname).write_text(
        f"Course: {cn}\nWeek: {wn}\nTopic: {tp}\n{'='*60}\n\n{notes}",
        encoding="utf-8")
    if log: log(f"✓ Notes saved → {fname}")
    return fname


def generate_slides(info, log=None):
    cn   = f"{info['course_name']} ({info['course_code']})"
    wn   = info["week_num"]; tp = info["topic"]; subs = info["subtopics"]
    s0   = subs[0] if subs else "subtopic1"
    s1   = subs[1] if len(subs) > 1 else "subtopic2"
    s2   = subs[2] if len(subs) > 2 else "subtopic3"
    if log: log("Waiting 6s before slide call (rate limit protection)…")
    time.sleep(6)
    if log: log(f"Gemini call 2 — 15 slides on {tp}…")
    raw = gemini_call(
        f"Create 15 slide contents for a university lecture.\n"
        f"Course: {cn} | Week: {wn} | Topic: {tp}\n"
        f"Subtopics: {', '.join(subs)} | Outcome: {info['outcome']}\n\n"
        f"Return ONLY a valid JSON array of 15 objects. No markdown. No backticks.\n"
        f"Each: {{\"slide\":N,\"type\":\"string\",\"title\":\"string\",\"content\":\"point1|point2|point3|point4\"}}\n"
        f"Use | between bullets. Max 8 words each. Real content about {tp}.\n"
        f"Types in order: 1=title,2=overview(5 bullets),3=intro,4={s0},5={s0} detail,"
        f"6={s1},7={s1} detail,8={s2},9=example(5 steps),10=diagram,11=algorithms,"
        f"12=applications,13=tips,14=summary(6 bullets),15=next\n"
        f"Return ONLY the JSON array.", 0.4)
    raw = re.sub(r"```json\s*|\s*```", "", raw).strip()
    try:
        data = json.loads(raw)
        if log: log(f"✓ {len(data)} slides received")
        return data
    except:
        m = re.search(r"\[.*\]", raw, re.DOTALL)
        if m:
            try:    data = json.loads(m.group())
            except: data = []
        else: data = []
        if log: log(f"⚠ JSON parse — got {len(data)} slides")
        return data


def build_pptx(slides, out, course, week, topic):
    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    NA = RGBColor(0x1e, 0x1b, 0x4b); IN = RGBColor(0x63, 0x66, 0xf1)
    BL = RGBColor(0x37, 0x8a, 0xdd); GO = RGBColor(0xea, 0xb3, 0x08)
    WH = RGBColor(0xff, 0xff, 0xff); LT = RGBColor(0xf8, 0xfa, 0xff)
    GR = RGBColor(0x6b, 0x72, 0x80); LA = RGBColor(0xa5, 0xb4, 0xfc)
    MI = RGBColor(0xee, 0xf2, 0xff); GN = RGBColor(0x05, 0x96, 0x69)
    RE = RGBColor(0xdc, 0x26, 0x26); DI = RGBColor(0xc7, 0xd2, 0xfe)
    ly = prs.slide_layouts[6]

    def R(sl, l, t, w, h, c):
        s = sl.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = c; s.line.fill.background()

    def T(sl, tx, l, t, w, h, sz=18, bold=False, c=None, al=PP_ALIGN.LEFT):
        tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True
        p  = tf.paragraphs[0]; p.alignment = al
        r  = p.add_run(); r.text = str(tx)
        r.font.size = Pt(sz); r.font.bold = bold
        if c: r.font.color.rgb = c

    def B(sl, items, l, t, w, h, sz=17, tc=None, dc=None):
        if not items: return
        tb = sl.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        tf = tb.text_frame; tf.word_wrap = True; first = True
        for item in items:
            item = str(item).strip()
            if not item: continue
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False; p.alignment = PP_ALIGN.LEFT; p.space_after = Pt(7)
            d = p.add_run(); d.text = "▸  "; d.font.size = Pt(sz - 2); d.font.color.rgb = dc or IN
            r = p.add_run(); r.text = item; r.font.size = Pt(sz); r.font.color.rgb = tc or NA

    for idx, s in enumerate(slides[:15]):
        sl  = prs.slides.add_slide(ly)
        st2 = s.get("type", "content")
        stl = s.get("title", f"Slide {idx+1}")
        items = [b.strip() for b in s.get("content", "").split("|") if b.strip()]
        sn    = idx + 1
        if sn == 1:
            R(sl, 0, 0, 13.33, 7.5, NA); R(sl, 0, 5.8, 13.33, 1.7, IN)
            R(sl, 0.6, 3.35, 5.5, 0.06, GO)
            for i, x in enumerate([11.5, 11.9, 12.3]):
                sh = sl.shapes.add_shape(9, Inches(x), Inches(0.35), Inches(0.25), Inches(0.25))
                sh.fill.solid(); sh.fill.fore_color.rgb = [IN, GO, BL][i]; sh.line.fill.background()
            T(sl, course, 0.6, 1.5, 12, 0.5, sz=15, c=LA)
            T(sl, topic,  0.6, 2.3, 11.5, 1.8, sz=38, bold=True, c=WH)
            T(sl, f"Week {week}  ·  Lecture Notes", 0.6, 5.95, 8, 0.5, sz=13, c=WH)
            T(sl, datetime.now(SL_TZ).strftime("%B %Y"), 9.5, 5.95, 3.5, 0.5, sz=13, c=WH, al=PP_ALIGN.RIGHT)
        else:
            R(sl, 0, 0, 13.33, 7.5, LT); R(sl, 0, 0, 13.33, 1.25, NA)
            R(sl, 0, 0, 0.1, 7.5, IN);   R(sl, 11.8, 0.38, 1.35, 0.42, IN)
            T(sl, f"{sn}/15", 11.82, 0.39, 1.3, 0.4, sz=11, bold=True, c=WH, al=PP_ALIGN.CENTER)
            T(sl, f"{course}  ·  Week {week}", 0.25, 0.06, 9, 0.38, sz=11, c=LA)
            T(sl, stl, 0.25, 0.42, 11.3, 0.7, sz=24, bold=True, c=WH)
            if st2 == "summary":
                R(sl, 0.25, 1.35, 12.8, 5.85, MI)
                B(sl, items, 0.55, 1.5, 12.3, 5.6, sz=20, tc=NA, dc=GN)
            elif st2 == "example":
                R(sl, 0.25, 1.35, 0.06, 5.85, GO)
                T(sl, "Worked Example", 0.45, 1.38, 6, 0.45, sz=13, c=GR)
                B(sl, items, 0.55, 1.9, 12.3, 5, sz=18, dc=GO)
            elif st2 == "tips":
                R(sl, 0.25, 1.35, 0.06, 5.85, RE)
                B(sl, items, 0.55, 1.5, 12.3, 5.6, sz=18, dc=RE)
            elif len(items) >= 6:
                mid = len(items) // 2
                B(sl, items[:mid], 0.25, 1.4, 6.2, 5.8, sz=17)
                R(sl, 6.6, 1.45, 0.03, 5.6, DI)
                B(sl, items[mid:], 6.7, 1.4, 6.3, 5.8, sz=17)
            else:
                B(sl, items, 0.35, 1.45, 12.6, 5.8, sz=19)
            R(sl, 0.1, 7.15, 13.2, 0.02, DI)

    while len(prs.slides) < 15:
        sl = prs.slides.add_slide(ly); n = len(prs.slides)
        R(sl, 0, 0, 13.33, 7.5, LT); R(sl, 0, 0, 13.33, 1.25, NA)
        R(sl, 0, 0, 0.1, 7.5, IN)
        T(sl, f"Slide {n}", 0.25, 0.42, 11.3, 0.7, sz=24, bold=True, c=WH)

    prs.save(str(out))


def run_phase2_pipeline(log=None):
    if log: log(" Step 1 — Parsing syllabus PDF…")
    info = parse_syllabus(SYLLABUS_PATH)
    cn   = f"{info['course_name']} ({info['course_code']})"
    wn   = info["week_num"]; tp = info["topic"]
    if log: log(f"   Week {wn}: {tp} | {', '.join(info['subtopics'])}")

    try:
        q = json.load(open("priority_queue.json"))
        if q:
            info["calendar_class"] = q[0].get("summary", "")
            info["class_time"]     = q[0].get("start", "")
            info["location"]       = q[0].get("location", "")
    except: pass

    if log: log(" Step 2 — Generating lecture notes…")
    nf  = generate_notes(info, log=log)

    if log: log(" Step 3 — Generating slide content…")
    sd  = generate_slides(info, log=log)

    sfn = f"slides_week{wn}_{tp.replace(' ','_')[:20]}_v1.pptx"
    sp  = OUTPUT_DIR / sfn
    if log: log(f" Step 4 — Building {sfn}…")
    build_pptx(sd, sp, cn, wn, tp)
    if log: log(f"   ✓ PPTX saved ({sp.stat().st_size // 1024} KB)")

    try:    nt = info["full_schedule"].get(str(wn + 1), {}).get("title", "TBD")
    except: nt = "TBD"
    save_memory(wn, cn, tp, info["subtopics"], info["outcome"], str(OUTPUT_DIR / nf), str(sp))
    if log: log(" Step 5 — Memory saved ✓")

    return {
        "course": cn, "week": wn, "topic": tp,
        "notes_fname": nf, "slides_fname": sfn,
        "next_topic": nt,
        "slides_path": str(sp),
        "notes_path":  str(OUTPUT_DIR / nf),
        "slide_count": len(sd),
    }


# ─────────────────────────────────────────────────────────────────────────────
#  COMBINED RUNNER
# ─────────────────────────────────────────────────────────────────────────────
def run_full_pipeline(log=None, to_email="csanduni123@gmail.com"):
    timeline = []
    start_ts = datetime.now(SL_TZ)

    # ── Phase 1 ──────────────────────────────────────────────────────────────
    if log: log(" **Phase 1** — Fetching calendar & analysing schedule…")
    p1_start  = datetime.now(SL_TZ)
    p1_result = run_phase1_agent()
    p1_end    = datetime.now(SL_TZ)
    timeline.append({
        "phase": "Phase 1 — Calendar & Schedule",
        "status": "complete",
        "time": p1_end.strftime("%H:%M:%S"),
        "duration": round((p1_end - p1_start).total_seconds(), 1),
        "detail": p1_result,
    })
    if log: log(f"   ✅ Phase 1 done at {p1_end.strftime('%H:%M:%S')}")

    # ── Load Phase 1 outputs ─────────────────────────────────────────────────
    try:    all_events = json.load(open("all_events.json"))
    except: all_events = []
    try:    queue = json.load(open("priority_queue.json"))
    except: queue = []

    # ── Phase 2 ──────────────────────────────────────────────────────────────
    if log: log(" **Phase 2** — Generating notes & slides…")
    p2_start  = datetime.now(SL_TZ)
    p2_result = run_phase2_pipeline(log=log)
    p2_end    = datetime.now(SL_TZ)
    timeline.append({
        "phase": "Phase 2 — Notes & Slides",
        "status": "complete",
        "time": p2_end.strftime("%H:%M:%S"),
        "duration": round((p2_end - p2_start).total_seconds(), 1),
        "detail": f"Week {p2_result['week']}: {p2_result['topic']} · {p2_result['slide_count']} slides",
    })
    if log: log(f"   ✅ Phase 2 done at {p2_end.strftime('%H:%M:%S')}")

    # ── Phase 3 ──────────────────────────────────────────────────────────────
    if log: log(" **Phase 3** — Emailing slides for lecturer review…")
    p3_start = datetime.now(SL_TZ)
    p3_result = act_send_email(to_email, log=log)
    p3_end = datetime.now(SL_TZ)
    timeline.append({
        "phase": "Phase 3 — Email Review",
        "status": "waiting",
        "time": p3_end.strftime("%H:%M:%S"),
        "duration": round((p3_end - p3_start).total_seconds(), 1),
        "detail": f"Sent {p3_result['pptx']} · tracking [{p3_result['tag']}]",
    })
    if log: log(f"Phase 3 — Started at {p3_end.strftime('%H:%M:%S')} — waiting for lecturer reply")

    total = round((datetime.now(SL_TZ) - start_ts).total_seconds(), 1)

    return {
        "all_events":  all_events,
        "queue":       queue,
        "p1_result":   p1_result,
        "p2_result":   p2_result,
        "p3_result":   p3_result,
        "timeline":    timeline,
        "total_secs":  total,
        "finished_at": datetime.now(SL_TZ).strftime("%H:%M:%S"),
    }



# ─────────────────────────────────────────────────────────────────────────────
#  STREAMLIT UI  — White + Blue/Yellow Gradient Professional Dashboard
# ─────────────────────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Lecturer AI Agent",
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="collapsed",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Plus+Jakarta+Sans:wght@300;400;500;600;700;800&family=JetBrains+Mono:wght@400;500&display=swap');

* { font-family: 'Plus Jakarta Sans', sans-serif; }

.stApp {
    background: linear-gradient(135deg, #ffffff 0%, #eff6ff 45%, #fefce8 80%, #f0f9ff 100%);
    color: #1e293b;
}
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding: 0 2.5rem 3rem 2.5rem; max-width: 1440px; }

.stButton > button {
    background: linear-gradient(135deg, #3b82f6 0%, #6366f1 60%, #f59e0b 100%) !important;
    color: white !important; border: none !important; border-radius: 10px !important;
    font-weight: 600 !important; font-size: 14px !important; padding: 12px 28px !important;
    box-shadow: 0 4px 16px rgba(59,130,246,0.3); transition: all 0.2s ease;
}
.stButton > button:hover { transform: translateY(-2px); box-shadow: 0 8px 28px rgba(59,130,246,0.45); }

div[data-testid="metric-container"] {
    background: white !important; border-radius: 14px !important;
    border: 0.5px solid #bfdbfe !important; padding: 16px !important;
    box-shadow: 0 1px 8px rgba(59,130,246,0.07) !important;
}
div[data-testid="metric-container"] label {
    color: #6b7280 !important; font-size: 10px !important; font-weight: 600 !important;
    letter-spacing: 1.2px !important; text-transform: uppercase !important;
}
div[data-testid="metric-container"] div[data-testid="metric-value"] {
    color: #1e3a8a !important; font-size: 32px !important; font-weight: 700 !important;
}
div[data-testid="metric-container"] div[data-testid="metric-delta"] { font-size: 11px !important; }

div[data-testid="stStatus"] {
    background: white !important; border: 0.5px solid #bfdbfe !important;
    border-radius: 14px !important; box-shadow: 0 2px 12px rgba(59,130,246,0.08) !important;
}
div[data-testid="stAlert"] { background: white !important; border-radius: 12px !important; }
details {
    background: white !important; border-radius: 12px !important;
    border: 0.5px solid #bfdbfe !important;
}

.card {
    background: white; border-radius: 14px; border: 0.5px solid #bfdbfe;
    padding: 14px 16px; margin-bottom: 8px;
    box-shadow: 0 1px 6px rgba(59,130,246,0.06);
}
.card-ql { border-left: 3px solid #6366f1; border-radius: 0 14px 14px 0; }
.card-mem { border-left: 3px solid #3b82f6; border-radius: 0 14px 14px 0; }
.card-status { border-left: 3px solid #16a34a; border-radius: 0 14px 14px 0; background: linear-gradient(135deg,#f0fdf4,#eff6ff); }

.tag { display:inline-block; padding:2px 8px; border-radius:20px; font-size:10px; font-weight:600; letter-spacing:0.4px; }
.tag-tmr  { background:#f5f3ff; color:#5b21b6; border:0.5px solid #c4b5fd; }
.tag-busy { background:#fff7ed; color:#c2410c; border:0.5px solid #fed7aa; }
.tag-done { background:#f0fdf4; color:#15803d; border:0.5px solid #bbf7d0; }
.tag-crit { background:#fef2f2; color:#dc2626; border:0.5px solid #fecaca; }
.tag-soon { background:#fffbeb; color:#d97706; border:0.5px solid #fde68a; }
.tag-plan { background:#eff6ff; color:#1d4ed8; border:0.5px solid #bfdbfe; }

.sec-title {
    font-size:10px; font-weight:700; color:#6b7280; letter-spacing:2px;
    text-transform:uppercase; margin:20px 0 8px 0; padding-bottom:6px;
    border-bottom:1.5px solid #dbeafe;
}
.mono { font-family:'JetBrains Mono',monospace; font-size:11px; color:#6b7280; }

.tl-row { display:flex; align-items:flex-start; gap:10px; padding:8px 0; border-bottom:0.5px solid #f1f5f9; }
.tl-dot  { width:8px; height:8px; border-radius:50%; background:#10b981; margin-top:5px; flex-shrink:0; }

.log-line { font-family:'JetBrains Mono',monospace; font-size:11.5px; padding:5px 0; border-bottom:0.5px solid #f1f5f9; line-height:1.6; display:flex; align-items:center; gap:8px; }
.log-ts   { color:#94a3b8; flex-shrink:0; }
.log-phase { color:#1d4ed8; font-weight:700; }
.log-done  { color:#15803d; }
.log-step  { color:#374151; }
.log-msg   { flex:1; }

.blink-dot {
    width:7px; height:7px; border-radius:50%; flex-shrink:0; display:inline-block;
}
.blink-dot.green  { background:#10b981; animation:bd-g 1.5s infinite; }
.blink-dot.blue   { background:#3b82f6; animation:bd-b 1.5s infinite; }
.blink-dot.gray   { background:#94a3b8; }
@keyframes bd-g { 0%,100%{opacity:1;box-shadow:0 0 0 0 rgba(16,185,129,0.6)} 50%{opacity:.7;box-shadow:0 0 0 4px rgba(16,185,129,0)} }
@keyframes bd-b { 0%,100%{opacity:1;box-shadow:0 0 0 0 rgba(59,130,246,0.6)} 50%{opacity:.7;box-shadow:0 0 0 4px rgba(59,130,246,0)} }

.pulse-g { width:7px; height:7px; border-radius:50%; background:#10b981; display:inline-block; margin-right:5px; vertical-align:middle; animation:pg 2s infinite; }
@keyframes pg { 0%,100%{opacity:1;box-shadow:0 0 0 0 rgba(16,185,129,0.5)} 50%{opacity:.8;box-shadow:0 0 0 4px rgba(16,185,129,0)} }
</style>
""", unsafe_allow_html=True)

# ── INIT ──────────────────────────────────────────────────────────────────────
now_sl = datetime.now(SL_TZ)
now_str = now_sl.strftime("%A %d %B %Y · %H:%M")

try:    syl = parse_syllabus(SYLLABUS_PATH)
except: syl = None

week_label   = f"Week {syl['week_num']} — {syl['topic']}"   if syl else "Syllabus not loaded"
course_label = f"{syl['course_name']} ({syl['course_code']})" if syl else ""
next_week_title = syl['full_schedule'].get(str(syl['week_num']+1),{}).get('title','TBD') if syl else "TBD"

# ── TOP BAR ───────────────────────────────────────────────────────────────────
st.markdown(f"""
<div style="background:linear-gradient(90deg,#1e3a8a 0%,#1d4ed8 60%,#b45309 100%);
  border-radius:14px;padding:10px 22px;display:flex;align-items:center;
  justify-content:space-between;margin-bottom:18px;flex-wrap:wrap;gap:8px;">
  <div style="display:flex;align-items:center;gap:10px;">
    <span style="font-size:14px;font-weight:600;color:#ffffff;letter-spacing:0.5px;">◈ Lecturer AI Agent</span>
    <span style="background:rgba(255,255,255,0.18);color:#e0f2fe;border:0.5px solid rgba(255,255,255,0.35);
      border-radius:20px;padding:3px 12px;font-size:11px;font-weight:500;">
      <span class='pulse-g'></span>Agent online
    </span>
  </div>
  <div style="display:flex;align-items:center;gap:16px;flex-wrap:wrap;">
    <span style="font-size:11px;color:#bfdbfe;">{now_str} (SL)</span>
    <span style="font-size:10px;color:rgba(191,219,254,0.7);">Auto-run · 07:00 SL daily</span>
  </div>
</div>
""", unsafe_allow_html=True)

# ── TITLE ─────────────────────────────────────────────────────────────────────
st.markdown("""
<div style="margin-bottom:14px;">
  <h1 style="font-size:26px;font-weight:700;color:#1e3a8a;margin:0 0 3px 0;letter-spacing:-0.3px;">
    Lecturer AI Agent
  </h1>
  <p style="font-size:12px;color:#6b7280;margin:0;">
    Calendar Check &nbsp;·&nbsp; Schedule Analysis &nbsp;·&nbsp; Lecture Notes &nbsp;·&nbsp; Slide Generation &nbsp;·&nbsp; Email Review &nbsp;·&nbsp; Canvas Upload
  </p>
</div>
""", unsafe_allow_html=True)

# ── KEY CHECK ─────────────────────────────────────────────────────────────────
missing_keys = check_keys()
if missing_keys:
    st.markdown(f"""
    <div class='card' style='border-left:3px solid #dc2626;border-radius:0 14px 14px 0;'>
      <div style='color:#dc2626;font-weight:700;margin-bottom:6px;'>⚠ Missing API Keys</div>
      <div style='color:#64748b;font-size:13px;'>
        Set <span class='mono'>{' and '.join(missing_keys)}</span> before running.<br><br>
        <b style='color:#1e293b;'>Windows:</b> <span class='mono'>set GEMINI_API_KEY=xxx &amp; set GOOGLE_CAL_KEY=xxx</span><br>
        <b style='color:#1e293b;'>Mac/Linux:</b> <span class='mono'>export GEMINI_API_KEY="xxx" &amp;&amp; export GOOGLE_CAL_KEY="xxx"</span>
      </div>
    </div>
    """, unsafe_allow_html=True)
    st.stop()

# ── SYLLABUS CHECK ────────────────────────────────────────────────────────────
syllabus_ok = Path(SYLLABUS_PATH).exists()
if not syllabus_ok:
    st.markdown("""
    <div class='card' style='border-left:3px solid #d97706;border-radius:0 14px 14px 0;margin-bottom:16px;'>
      <div style='color:#d97706;font-weight:700;margin-bottom:4px;'>📄 Syllabus PDF Required</div>
      <div style='color:#64748b;font-size:13px;'>Place <b>syllabus.pdf</b> in the working directory or upload below.</div>
    </div>
    """, unsafe_allow_html=True)
    up = st.file_uploader("Upload syllabus.pdf", type=["pdf"], label_visibility="collapsed")
    if up:
        with open(SYLLABUS_PATH, "wb") as fh: fh.write(up.read())
        st.success("✅ syllabus.pdf saved. Reload the page.")
        st.stop()
    st.stop()

# ── RUN BUTTON ────────────────────────────────────────────────────────────────
col_btn, col_email, col_desc = st.columns([1, 1.35, 2.65])
with col_btn:
    run_btn = st.button("⚡  Run Autonomous Agent", use_container_width=True)
with col_email:
    lecturer_email = st.text_input("Lecturer email", "csanduni123@gmail.com", label_visibility="collapsed")
with col_desc:
    st.markdown("""
    <p style='color:#6b7280;font-size:13px;padding-top:10px;'>
      Runs <b style='color:#1d4ed8;'>Phase 1</b> (calendar → queue),
      <b style='color:#6366f1;'>Phase 2</b> (notes → slides), then
      <b style='color:#b45309;'>Phase 3</b> (email review → auto-poll), then Canvas upload after approval — fully autonomous.
    </p>""", unsafe_allow_html=True)

# ── BADGE MAP ─────────────────────────────────────────────────────────────────
BADGE = {
    "ongoing":  ("🔴", "tag-crit"),
    "critical": ("🔴", "tag-crit"),
    "soon":     ("🟡", "tag-soon"),
    "planned":  ("🔵", "tag-plan"),
    "tomorrow": ("🟣", "tag-tmr"),
}
ACCENT = {
    "critical": "card-ql", "soon": "card-ql",
    "planned": "card-ql",  "tomorrow": "card-ql",
}

# ── RUN PIPELINE ──────────────────────────────────────────────────────────────
def is_relevant_live_log(msg: str) -> bool:
    text = str(msg).strip()
    if not text:
        return False
    noise = [
        "Gemini call", "Waiting 6s", "rate limit", "Subject:", "From:",
        "skipping", "Found ", "Polling [", "No lecturer reply yet",
        "Upload HTTP status", "Course:", "Folder:", "File:", "Name:",
        "Calling Gemini", "Summary:", "Think:", "Authenticating Gmail",
    ]
    if any(n in text for n in noise):
        return False
    keep = [
        "Phase 1", "Phase 2", "Phase 3", "Phase 4",
        "Step 1", "Step 2", "Step 3", "Step 4", "Step 5",
        "Notes saved", "slides received", "PPTX saved", "Memory saved",
        "Email sent", "waiting for lecturer reply", "reply found",
        "Changes detected", "Revision", "APPROVED", "Approved file saved",
        "Canvas upload", "Canvas announcement", "students notified",
        "complete", "saved", "sent", "uploaded", "Found approved file",
    ]
    return any(k.lower() in text.lower() for k in keep)

if run_btn:
    log_lines = []

    def live_log(msg):
        ts = datetime.now(SL_TZ).strftime("%H:%M:%S")
        if not is_relevant_live_log(msg):
            return
        log_lines.append({"ts": ts, "msg": msg})
        st.write(f"`{ts}` {msg}")

    with st.status("⚡ Autonomous Agent Running…", expanded=True) as sw:
        try:
            data = run_full_pipeline(log=live_log, to_email=lecturer_email)
            sw.update(label="✅ Pipeline complete", state="complete")
            st.session_state["run_data"]  = data
            st.session_state["ran_at"]    = datetime.now(SL_TZ).strftime("%H:%M:%S")
            st.session_state["log_lines"] = log_lines
        except Exception as e:
            sw.update(label="❌ Pipeline failed", state="error")
            st.error(f"Error: {e}")
            st.code(traceback.format_exc())
            st.stop()

# ── RESULTS ───────────────────────────────────────────────────────────────────
if "run_data" in st.session_state:
    d          = st.session_state["run_data"]
    all_events = d["all_events"]
    queue      = d["queue"]
    p2         = d["p2_result"]
    timeline   = d["timeline"]
    log_lines  = st.session_state.get("log_lines", [])

    today_ev      = [e for e in all_events if e["day_label"] == "today"]
    tomorrow_ev   = [e for e in all_events if e["day_label"] == "tomorrow"]
    critical_n    = sum(1 for e in queue if e["urgency"] == "critical")
    today_busy    = len(today_ev) >= 2
    tomorrow_busy = len(tomorrow_ev) >= 2

    # ── Completion banner ─────────────────────────────────────────────────────
    st.markdown(f"""
    <div style='background:linear-gradient(135deg,#f0fdf4,#eff6ff);border:0.5px solid #bbf7d0;
      border-radius:14px;padding:12px 20px;margin:14px 0;
      display:flex;align-items:center;justify-content:space-between;flex-wrap:wrap;gap:8px;
      box-shadow:0 1px 8px rgba(16,185,129,0.08);'>
      <span style='font-size:13px;font-weight:600;color:#15803d;'>
        ✅ Four-step workflow started &nbsp;·&nbsp; Total: {d['total_secs']}s &nbsp;·&nbsp; Finished {d['finished_at']} SL
      </span>
      <span class='mono'>Phase 1: {timeline[0]['duration']}s &nbsp;·&nbsp; Phase 2: {timeline[1]['duration']}s &nbsp;·&nbsp; Phase 3: {timeline[2]['duration']}s &nbsp;·&nbsp; Phase 4: after approval</span>
    </div>
    """, unsafe_allow_html=True)

    # ── KPI METRICS ───────────────────────────────────────────────────────────
    st.markdown("<div class='sec-title'>Schedule Overview</div>", unsafe_allow_html=True)

    # next class time helpers
    def get_next_time(events):
        upcoming = [e for e in events if e.get("urgency") != "ongoing"]
        return upcoming[0]["start"][11:] if upcoming else "—"

    today_next    = get_next_time(today_ev)
    tomorrow_next = get_next_time(tomorrow_ev)

    today_delta_color    = "#c2410c" if today_busy else "#15803d"
    today_delta_bg       = "#fff7ed" if today_busy else "#f0fdf4"
    today_delta_border   = "#fed7aa" if today_busy else "#bbf7d0"
    today_delta_txt      = "Busy day" if today_busy else "Normal load"

    tmr_delta_color      = "#c2410c" if tomorrow_busy else "#15803d"
    tmr_delta_bg         = "#fff7ed" if tomorrow_busy else "#f0fdf4"
    tmr_delta_border     = "#fed7aa" if tomorrow_busy else "#bbf7d0"
    tmr_delta_txt        = "Busy day" if tomorrow_busy else "Normal load"

    q_delta_color        = "#dc2626" if critical_n else "#15803d"
    q_delta_bg           = "#fef2f2" if critical_n else "#f0fdf4"
    q_delta_border       = "#fecaca" if critical_n else "#bbf7d0"
    q_delta_txt          = f"{critical_n} critical" if critical_n else "On track"

    m1, m2, m3, m4 = st.columns(4)

    with m1:
        st.markdown(f"""
        <div style='background:white;border:0.5px solid #bfdbfe;border-radius:14px;padding:16px;
          box-shadow:0 1px 8px rgba(59,130,246,0.07);'>
          <div style='font-size:10px;font-weight:700;color:#6b7280;letter-spacing:1.2px;text-transform:uppercase;margin-bottom:6px;'>Classes Today</div>
          <div style='font-size:34px;font-weight:700;color:#1e3a8a;line-height:1;margin-bottom:6px;'>{len(today_ev)}</div>
          <div style='margin-bottom:5px;'>
            <span style='background:{today_delta_bg};color:{today_delta_color};border:0.5px solid {today_delta_border};
              border-radius:20px;padding:2px 8px;font-size:10px;font-weight:600;'>
              {"🔥 " if today_busy else ""}{today_delta_txt}
            </span>
          </div>
          <div style='font-size:10px;color:#94a3b8;font-family:monospace;'>Next: {today_next}</div>
        </div>""", unsafe_allow_html=True)

    with m2:
        st.markdown(f"""
        <div style='background:white;border:0.5px solid #bfdbfe;border-radius:14px;padding:16px;
          box-shadow:0 1px 8px rgba(59,130,246,0.07);'>
          <div style='font-size:10px;font-weight:700;color:#6b7280;letter-spacing:1.2px;text-transform:uppercase;margin-bottom:6px;'>Classes Tomorrow</div>
          <div style='font-size:34px;font-weight:700;color:#1e3a8a;line-height:1;margin-bottom:6px;'>{len(tomorrow_ev)}</div>
          <div style='margin-bottom:5px;'>
            <span style='background:{tmr_delta_bg};color:{tmr_delta_color};border:0.5px solid {tmr_delta_border};
              border-radius:20px;padding:2px 8px;font-size:10px;font-weight:600;'>
              {"🔥 " if tomorrow_busy else ""}{tmr_delta_txt}
            </span>
          </div>
          <div style='font-size:10px;color:#94a3b8;font-family:monospace;'>Starts: {tomorrow_next}</div>
        </div>""", unsafe_allow_html=True)

    with m3:
        st.markdown(f"""
        <div style='background:white;border:0.5px solid #bfdbfe;border-radius:14px;padding:16px;
          box-shadow:0 1px 8px rgba(59,130,246,0.07);'>
          <div style='font-size:10px;font-weight:700;color:#6b7280;letter-spacing:1.2px;text-transform:uppercase;margin-bottom:6px;'>Priority Queue</div>
          <div style='font-size:34px;font-weight:700;color:#1e3a8a;line-height:1;margin-bottom:6px;'>{len(queue)}</div>
          <div style='margin-bottom:5px;'>
            <span style='background:{q_delta_bg};color:{q_delta_color};border:0.5px solid {q_delta_border};
              border-radius:20px;padding:2px 8px;font-size:10px;font-weight:600;'>
              {"⚠ " if critical_n else ""}{q_delta_txt}
            </span>
          </div>
          <div style='font-size:10px;color:#94a3b8;font-family:monospace;'>priority_queue.json</div>
        </div>""", unsafe_allow_html=True)

    with m4:
        st.markdown(f"""
        <div style='background:linear-gradient(135deg,#eff6ff,#e0e7ff);border:0.5px solid #818cf8;
          border-radius:14px;padding:16px;box-shadow:0 1px 8px rgba(99,102,241,0.1);'>
          <div style='font-size:10px;font-weight:700;color:#4338ca;letter-spacing:1.2px;
            text-transform:uppercase;margin-bottom:6px;'>Current Week</div>
          <div style='font-size:15px;font-weight:700;color:#1e3a8a;line-height:1.2;margin-bottom:4px;'>
            {week_label}
          </div>
          <div style='font-size:11px;color:#6366f1;'>{course_label}</div>
        </div>
        """, unsafe_allow_html=True)

    st.markdown("<div style='height:6px;'></div>", unsafe_allow_html=True)

    # ── TWO COLUMN LAYOUT ─────────────────────────────────────────────────────
    col_left, col_right = st.columns([4, 5], gap="large")

    with col_left:

        # Priority Queue
        st.markdown("<div class='sec-title'>📋 Slide Preparation Priority Queue</div>", unsafe_allow_html=True)
        st.markdown("<div style='font-size:10px;color:#94a3b8;margin:-4px 0 10px 0;'>Ongoing classes excluded — in-advance only.</div>", unsafe_allow_html=True)

        if not queue:
            st.markdown("<div class='card' style='text-align:center;color:#94a3b8;font-size:13px;padding:20px;'>No lectures in queue.</div>", unsafe_allow_html=True)
        else:
            for i, e in enumerate(queue, 1):
                urg       = e.get("urgency", "planned")
                icon, cls = BADGE.get(urg, ("⚪", "tag-plan"))
                hrs       = e["hours_until"]
                hrs_str   = f"{hrs:.1f} hrs" if hrs > 0 else "Now"
                busy_badge = "<span class='tag tag-busy' style='margin-left:4px;'>🔥 Busy</span>" if e.get("busy_day") else ""
                action_txt = e.get("action","").replace("_"," ")
                summary_safe = html.escape(str(e.get("summary", "")))
                location_safe = html.escape(str(e.get("location", "")))
                action_safe = html.escape(str(action_txt))
                st.markdown(f"""
<div class="card card-ql" style="margin-bottom:8px;">
  <div style="display:flex;align-items:center;justify-content:space-between;margin-bottom:6px;flex-wrap:wrap;gap:8px;">
    <div style="display:flex;align-items:center;gap:5px;flex-wrap:wrap;">
      <span style="font-size:10px;color:#6366f1;font-weight:700;">#{i}</span>
      <span style="font-size:13px;font-weight:600;color:#0f172a;">{summary_safe}</span>
      <span class="tag {cls}">{icon} {urg.upper()}</span>
      {busy_badge}
    </div>
    <span class="mono">⏳ {hrs_str}</span>
  </div>
  <div class="mono" style="margin-top:6px;">
    🕐 {e['start'][11:]}–{e['end'][11:]} &nbsp;|&nbsp; 📍 {location_safe} &nbsp;|&nbsp; 📅 {e['day_label'].capitalize()} &nbsp;|&nbsp; → {action_safe}
  </div>
</div>
""", unsafe_allow_html=True)

        # Agent Activity Log
        st.markdown("<div class='sec-title'>🕐 Agent Activity Log</div>", unsafe_allow_html=True)
        st.markdown("<div class='card' style='padding:12px 16px;'>", unsafe_allow_html=True)
        activity_items = list(timeline)
        p4_status_activity = load_phase4_status()
        if p4_status_activity.get("phase") == "complete":
            try:
                p4_time = datetime.fromisoformat(p4_status_activity.get("uploaded_at", "")).strftime("%H:%M:%S")
            except Exception:
                p4_time = "—"
            activity_items.append({
                "phase": "Phase 4 — Canvas Upload",
                "status": "complete",
                "time": p4_time,
                "duration": "—",
                "detail": f"Uploaded to Canvas · folder {p4_status_activity.get('canvas_folder','—')} · students notified",
            })
        for item in activity_items:
            item_status = item.get("status", "complete")
            status_txt = "✓ Done" if item_status == "complete" else ("⏳ Waiting" if item_status == "waiting" else item_status.title())
            status_cls = "tag-done" if item_status == "complete" else "tag-plan"
            st.markdown(f"""
            <div class='tl-row'>
              <div class='tl-dot'></div>
              <div style='flex:1;'>
                <div style='font-size:12px;font-weight:600;color:#0f172a;'>{item['phase']}</div>
                <div class='mono' style='margin-top:2px;'>{item['detail']}</div>
              </div>
              <div style='text-align:right;flex-shrink:0;'>
                <span class='tag {status_cls}'>{status_txt}</span>
                <div class='mono' style='margin-top:3px;'>{item['time']} · {item['duration']}s</div>
              </div>
            </div>""", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)

        # Memory
        mem = load_memory()
        if mem:
            st.markdown("<div class='sec-title'>🧠 Agent Memory</div>", unsafe_allow_html=True)
            for k, v in list(mem.items())[-5:]:
                st.markdown(f"""
                <div class='card card-mem' style='padding:10px 14px;margin-bottom:6px;'>
                  <div style='font-size:12px;font-weight:600;color:#0f172a;'>Week {v['week']} · {v['course']}</div>
                  <div style='font-size:11px;color:#475569;margin-top:1px;'>{v['topic']}</div>
                  <div class='mono' style='margin-top:3px;'>{v['generated']}</div>
                </div>""", unsafe_allow_html=True)

    with col_right:

        # Live Agent Log (bigger panel)
        st.markdown("<div class='sec-title'>⚡ Live Agent Log</div>", unsafe_allow_html=True)

        log_html = ""
        import re as _re
        def strip_emoji(text):
            return _re.sub(r'[\U00010000-\U0010ffff\U00002600-\U000027BF\U0001F300-\U0001FAFF]', '', text).strip()

        # UI only: surface Phase 3 / Phase 4 completion in the main Live Agent Log.
        p3_status_for_log = load_status()
        p4_status_for_log = load_phase4_status()
        if p3_status_for_log.get("phase") == "approved":
            approved_at_raw = p3_status_for_log.get("approved_at", "")
            try:
                approved_time = datetime.fromisoformat(approved_at_raw).strftime("%H:%M:%S")
            except Exception:
                approved_time = datetime.now(SL_TZ).strftime("%H:%M:%S")
            if not any("Phase 3 done" in str(x.get("msg", "")) for x in log_lines):
                log_lines.append({"ts": approved_time, "msg": f"Phase 3 done at {approved_time} — APPROVED"})
        if p4_status_for_log.get("phase") == "complete":
            uploaded_at_raw = p4_status_for_log.get("uploaded_at", "")
            try:
                uploaded_time = datetime.fromisoformat(uploaded_at_raw).strftime("%H:%M:%S")
            except Exception:
                uploaded_time = datetime.now(SL_TZ).strftime("%H:%M:%S")
            if not any("Phase 4 done" in str(x.get("msg", "")) for x in log_lines):
                log_lines.append({"ts": uploaded_time, "msg": f"Phase 4 done at {uploaded_time} — students notified"})
        st.session_state["log_lines"] = log_lines

        for ll in log_lines:
            msg = strip_emoji(ll["msg"].strip())
            if not msg: continue
            msg_lower = msg.lower()
            if "done" in msg_lower or "approved" in msg_lower or "saved" in msg_lower or "complete" in msg_lower or "received" in msg_lower or "revision" in msg_lower or "upload" in msg_lower or "notified" in msg_lower or msg.startswith("✓"):
                dot = "<span class='blink-dot green'></span>"
                cls = "log-done"
                msg = msg.lstrip("✓ ").strip()
                msg = "✓ " + msg
            elif "Phase 1" in msg or "Phase 2" in msg or "Phase 3" in msg or "Phase 4" in msg:
                dot = "<span class='blink-dot blue'></span>"
                cls = "log-phase"
            else:
                dot = "<span class='blink-dot gray'></span>"
                cls = "log-step"
            log_html += f"<div class='log-line {cls}'>{dot}<span class='log-ts'>{ll['ts']}</span><span class='log-msg'>{msg}</span></div>"

        if not log_html:
            log_html = "<div class='log-line log-step'><span class='blink-dot gray'></span><span class='log-msg'>No log yet — run the agent first.</span></div>"

        st.markdown(f"""
        <div class='card' style='padding:14px 18px;margin-bottom:12px;'>
          {log_html}
        </div>""", unsafe_allow_html=True)

        # Current Status + Phase 3 review state
        st.markdown("<div class='sec-title'>Current Status</div>", unsafe_allow_html=True)
        p3_status = load_status()
        p3_phase = p3_status.get("phase", "waiting") if p3_status else "waiting"
        p3_tag = p3_status.get("subject_tag", d.get("p3_result", {}).get("tag", "—")) if p3_status else d.get("p3_result", {}).get("tag", "—")
        p3_rev = p3_status.get("revision_count", 0) if p3_status else 0
        p3_reply = p3_status.get("reply_body", "") if p3_status else ""
        p4_status = load_phase4_status()
        p4_phase = p4_status.get("phase", "waiting")
        phase_label = {
            "waiting": "Phase 3 active — waiting for lecturer reply",
            "approved": "Approved — Canvas upload active",
            "changes": "Changes detected — revision loop active",
            "unclear": "Reply received — needs clearer approval or changes",
        }.get(p3_phase, "Phase 3 active")
        if p4_phase == "complete":
            phase_label = "All phases complete — slides uploaded and students notified"
        phase_color = "#15803d" if p4_phase == "complete" or p3_phase == "approved" else ("#b45309" if p3_phase == "waiting" else "#1d4ed8")
        st.markdown(f"""
        <div class='card card-status' style='padding:14px 18px;'>
          <div style='font-size:13px;font-weight:700;color:{phase_color};margin-bottom:10px;'>
            ✓ Pipeline complete — {phase_label}
          </div>
          <div style='display:flex;flex-direction:column;gap:5px;'>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Course</span><span style='color:#374151;'>{p2['course']}</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Week</span><span style='color:#374151;'>{p2['week']} — {p2['topic']}</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Slides</span><span style='color:#374151;'>{p2['slide_count']} generated ✓</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Review</span><span style='color:#374151;'>[{p3_tag}] · revision #{p3_rev}</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Phase 3</span><span style='color:#374151;'>{p3_phase} ✓</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Canvas</span><span style='color:#374151;'>{p4_phase} {"✓" if p4_phase == "complete" else "— after approval"}</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Queue</span><span style='color:#374151;'>priority_queue.json ✓</span></div>
            <div class='mono'><span style='color:#4338ca;font-weight:600;min-width:74px;display:inline-block;'>Output</span><span style='color:#374151;'>output_slides/ ✓</span></div>
          </div>
        </div>
        """, unsafe_allow_html=True)

        # If Phase 3 is already approved, continue Phase 4 inside existing workflow/status.
        if p3_status and p3_phase == "approved" and load_phase4_status().get("phase") != "complete" and not st.session_state.get("phase4_attempted", False):
            st.session_state["phase4_attempted"] = True
            def p4log(m):
                if not is_relevant_live_log(m):
                    return
                ts = datetime.now(SL_TZ).strftime('%H:%M:%S')
                live_logs = st.session_state.get("log_lines", [])
                live_logs.append({"ts": ts, "msg": m})
                st.session_state["log_lines"] = live_logs
            try:
                p4log("Phase 4 — Uploading approved slides to Canvas")
                p4_result = run_phase4_once(log=p4log)
                st.session_state["p4_result"] = p4_result
                p4log(f"Phase 4 done at {datetime.now(SL_TZ).strftime('%H:%M:%S')} — students notified")
                st.rerun()
            except Exception as e:
                p4log(f"Phase 4 waiting — {e}")

        # Phase 3 auto-poll embedded into the existing dashboard area.
        if p3_status and p3_phase == "waiting":
            st.markdown("<div class='sec-title'>📧 Email Review Auto-Poll</div>", unsafe_allow_html=True)
            if st.button("🔍 Poll Review Now", use_container_width=True):
                st.session_state["phase3_poll_now"] = True
                st.rerun()
            poll_due = st.session_state.get("phase3_poll_now", False) or st.session_state.get("phase3_auto_poll", True)
            if poll_due:
                st.session_state["phase3_poll_now"] = False
                p3_logs = []
                def p3log(m):
                    if not is_relevant_live_log(m):
                        return
                    ts = datetime.now(SL_TZ).strftime('%H:%M:%S')
                    p3_logs.append(f"{ts} — {m}")
                    live_logs = st.session_state.get("log_lines", [])
                    live_logs.append({"ts": ts, "msg": m})
                    st.session_state["log_lines"] = live_logs
                try:
                    poll_res = observe_poll(log=p3log)
                    pr = poll_res.get("result", "")
                    if pr in ("changes", "unclear"):
                        p3log("Changes detected — starting autonomous revision…")
                        rev_res = act_regenerate_and_resend(log=p3log)
                        p3log(f"Revision {rev_res['revision']} complete — {rev_res['new_pptx']} sent")
                        st.session_state["phase3_auto_poll"] = True
                        st.rerun()
                    elif pr == "approved":
                        done_time = datetime.now(SL_TZ).strftime('%H:%M:%S')
                        p3log(f"Phase 3 done at {done_time} — APPROVED")
                        st.session_state["phase3_auto_poll"] = False
                        try:
                            p3log("Phase 4 — Uploading approved slides to Canvas")
                            p4_result = run_phase4_once(log=p3log)
                            st.session_state["p4_result"] = p4_result
                            p3log(f"Phase 4 done at {datetime.now(SL_TZ).strftime('%H:%M:%S')} — students notified")
                            st.success("✅ Approved slides uploaded to Canvas and students notified.")
                        except Exception as e:
                            p3log(f"Phase 4 waiting — {e}")
                            st.warning(f"Phase 3 approved, but Canvas upload did not complete: {e}")
                        st.rerun()
                    else:
                        st.session_state["phase3_auto_poll"] = True
                    if p3_logs:
                        st.markdown("<div class='card' style='font-family:monospace;font-size:11px;line-height:1.7;'>" + "<br>".join(p3_logs[-8:]) + "</div>", unsafe_allow_html=True)
                    if st.session_state.get("phase3_auto_poll", True):
                        time.sleep(POLL_MINS * 60)
                        st.rerun()
                except Exception as e:
                    st.error(f"Phase 3 poll failed: {e}")
                    st.code(traceback.format_exc())
        elif p3_status and p3_reply:
            reply_time = "—"
            for time_key in ("approved_at", "last_polled", "sent_at"):
                if p3_status.get(time_key):
                    try:
                        reply_time = datetime.fromisoformat(p3_status[time_key]).strftime("%H:%M:%S")
                        break
                    except Exception:
                        pass
            reply_safe = html.escape(str(p3_reply[:500]))
            st.markdown(f"""
<div class="card" style="padding:12px 16px;">
  <div style="font-size:12px;font-weight:700;color:#1e3a8a;margin-bottom:6px;">Lecturer reply <span class="mono">{reply_time}</span></div>
  <div style="font-size:12px;color:#374151;">{reply_safe}</div>
</div>
""", unsafe_allow_html=True)

    # ── Footer ────────────────────────────────────────────────────────────────
    st.markdown(f"""
    <div style='background:linear-gradient(90deg,#eff6ff,#fefce8);border:0.5px solid #bfdbfe;
      border-radius:12px;padding:10px 20px;margin-top:16px;
      display:flex;align-items:center;justify-content:space-between;flex-wrap:wrap;gap:6px;'>
      <span style='font-size:12px;font-weight:600;color:#1e3a8a;'>
        ✅ Phase 1 + Phase 2 + Phase 3 + Canvas upload workflow started.
        <span style='font-weight:400;color:#6b7280;margin-left:6px;'>
          Files saved to
          <code style='background:#dbeafe;color:#1d4ed8;padding:1px 6px;border-radius:4px;font-size:10px;'>output_slides/</code>
          — Phase 4 uploads to Canvas after lecturer approval.
        </span>
      </span>
      <span class='mono'>{d['finished_at']} SL</span>
    </div>
    """, unsafe_allow_html=True)

else:
    # ── IDLE ──────────────────────────────────────────────────────────────────
    st.markdown("""
    <div style='background:white;border:0.5px solid #bfdbfe;border-radius:18px;
      padding:56px 40px;text-align:center;margin-top:20px;
      box-shadow:0 4px 20px rgba(59,130,246,0.06);'>
      <div style='font-size:40px;margin-bottom:12px;'>⚡</div>
      <h3 style='color:#1e3a8a;font-size:20px;font-weight:700;margin-bottom:8px;'>Agent Standby</h3>
      <p style='color:#64748b;font-size:13px;max-width:480px;margin:0 auto;line-height:1.8;'>
        Press <b style='color:#1d4ed8;'>Run Autonomous Agent</b> above to start the full pipeline.<br>
        Phase 1, Phase 2, Phase 3 and Canvas upload run automatically — no manual steps required.
      </p>
      <div style='margin-top:24px;display:flex;gap:8px;justify-content:center;flex-wrap:wrap;'>
        <span style='background:#eff6ff;color:#1d4ed8;border:0.5px solid #bfdbfe;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>🗓 Calendar fetch</span>
        <span style='color:#cbd5e1;font-size:13px;padding:5px 0;'>→</span>
        <span style='background:#eff6ff;color:#1d4ed8;border:0.5px solid #bfdbfe;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>📊 Schedule analysis</span>
        <span style='color:#cbd5e1;font-size:13px;padding:5px 0;'>→</span>
        <span style='background:#f5f3ff;color:#6d28d9;border:0.5px solid #ddd6fe;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>📝 Lecture notes</span>
        <span style='color:#cbd5e1;font-size:13px;padding:5px 0;'>→</span>
        <span style='background:#fffbeb;color:#b45309;border:0.5px solid #fde68a;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>🎨 15 slides</span>
        <span style='color:#cbd5e1;font-size:13px;padding:5px 0;'>→</span>
        <span style='background:#f0fdf4;color:#15803d;border:0.5px solid #bbf7d0;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>📧 Email review</span>
        <span style='color:#cbd5e1;font-size:13px;padding:5px 0;'>→</span>
        <span style='background:#eef2ff;color:#4338ca;border:0.5px solid #c7d2fe;border-radius:20px;padding:6px 14px;font-size:11px;font-weight:600;'>📤 Canvas upload</span>
      </div>
    </div>
    """, unsafe_allow_html=True)
